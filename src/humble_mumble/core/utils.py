"""
utils.py
========
Shared helper functions for date parsing, KPI cards, and Excel utilities.
Centralizing these removes duplicate implementations across report builders.
"""

from datetime import datetime, date
import re
from openpyxl.styles import Font, Border, Side
from humble_mumble.core import config

GENERIC_SERIALS = {
    "N/A", "NONE", "UNKNOWN", "", "TO BE FILLED BY O.E.M.", "SYSTEM SERIAL NUMBER",
    "DEFAULT STRING", "0", "00000000", "0123456789", "123456789", "NOT APPLICABLE",
    "CHASSIS SERIAL NUMBER", "TYPE2 - BOARD SERIAL NUMBER", "ALL", "DEFAULT",
    "SYSTEM MANUFACTURER", "INVALID", "TBD"
}

def is_valid_serial(serial):
    """Checks if a serial number is valid and not a placeholder OEM string."""
    if not serial:
        return False
    s_clean = str(serial).strip().upper()
    return s_clean not in GENERIC_SERIALS and len(s_clean) > 2

def normalize_hostname(name):
    """
    Normalizes a computer or agent hostname by stripping whitespace,
    converting to upper-case, and stripping any domain suffix (e.g., .corp.local, .domain.com).
    """
    if not name:
        return ""
    clean = str(name).strip().upper()
    if "." in clean:
        clean = clean.split(".")[0].strip()
    return clean

def safe_float(val, default=None):
    """Safely extracts a float from strings with units (e.g., '15.8 GiB', '16 GB', '75%')."""
    if val is None or val == "":
        return default
    if isinstance(val, (int, float)):
        return float(val)
    try:
        m = re.search(r'[-+]?\d*\.?\d+', str(val))
        if m:
            return float(m.group(0))
        return default
    except (ValueError, TypeError):
        return default

def safe_int(val, default=None):
    """Safely extracts an integer from values or strings."""
    if val is None or val == "":
        return default
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        return int(val)
    try:
        m = re.search(r'[-+]?\d+', str(val).strip())
        if m:
            return int(m.group(0))
        return default
    except (ValueError, TypeError):
        return default

def parse_date(s):
    """
    Parses a date string in various formats (ISO, YYYY-MM-DD, M/D/YYYY)
    into a datetime.date object. Returns None if unparseable or blank.
    """
    if not s:
        return None
    if isinstance(s, datetime):
        return s.date()
    if isinstance(s, date):
        return s
    if not isinstance(s, str):
        return None
    s = s.strip()
    if not s:
        return None

    # Handle ISO or YYYY-MM-DD (e.g. 2026-10-16 or 2026-10-16T10:30:00)
    if '-' in s:
        try:
            clean_s = s.replace('T', ' ').split()[0]
            parts = clean_s.split('-')
            if len(parts) == 3:
                yr = int(parts[0])
                if yr < 100:
                    yr += 2000
                return datetime(yr, int(parts[1]), int(parts[2])).date()
        except (ValueError, TypeError, IndexError):
            pass

    # Handle M/D/YYYY (e.g. 10/16/2026 or 10/16/2026, 10:30 AM)
    if '/' in s:
        try:
            date_part = s.split(',')[0].strip().split()[0]
            parts = date_part.split('/')
            if len(parts) == 3:
                yr = int(parts[2])
                if yr < 100:
                    yr += 2000
                return datetime(yr, int(parts[0]), int(parts[1])).date()
        except (ValueError, TypeError, IndexError):
            pass

    return None

# Backward compatibility alias
parse_ninja_date = parse_date

def create_kpi_card(ws, start_col, start_row, *args, **kwargs):
    """
    Creates a unified 2-row styled KPI card.
    Supports multiple calling signatures:
      - (ws, start_col, start_row, label, value, is_alert=False, label_on_top=False)
      - (ws, c1, r1, c2, r2, label, value, is_alert=False)
    """
    if len(args) >= 4:
        c1 = start_col
        r1 = start_row
        c2 = args[0]
        r2 = args[1]
        label = args[2]
        value = args[3]
        is_alert = args[4] if len(args) > 4 else kwargs.get("is_alert", False)
        label_on_top = kwargs.get("label_on_top", True)
    elif len(args) >= 2:
        c1 = start_col
        r1 = start_row
        c2 = start_col + kwargs.get("col_span", 2) - 1
        r2 = start_row + 1
        label = args[0]
        value = args[1]
        is_alert = args[2] if len(args) > 2 else kwargs.get("is_alert", False)
        label_on_top = args[3] if len(args) > 3 else kwargs.get("label_on_top", False)
    else:
        label = kwargs.get("label", "")
        value = kwargs.get("value", "")
        c1 = start_col
        r1 = start_row
        c2 = start_col + kwargs.get("col_span", 2) - 1
        r2 = start_row + 1
        is_alert = kwargs.get("is_alert", False)
        label_on_top = kwargs.get("label_on_top", False)

    ws.merge_cells(start_row=r1, start_column=c1, end_row=r1, end_column=c2)
    ws.merge_cells(start_row=r2, start_column=c1, end_row=r2, end_column=c2)

    val_font = Font(name="Calibri", size=18, bold=True, color=config.COLOR_ALERT_RED_FONT if is_alert else config.COLOR_PRIMARY_NAVY)

    if label_on_top:
        top_val, top_font = label, config.FONT_CARD_LABEL
        bot_val, bot_font = value, val_font
    else:
        top_val, top_font = value, val_font
        bot_val, bot_font = label, config.FONT_CARD_LABEL

    c_top = ws.cell(row=r1, column=c1, value=top_val)
    c_top.font = top_font
    c_top.alignment = config.ALIGN_CENTER
    c_top.fill = config.FILL_KPI_CARD

    c_bot = ws.cell(row=r2, column=c1, value=bot_val)
    c_bot.font = bot_font
    c_bot.alignment = config.ALIGN_CENTER
    c_bot.fill = config.FILL_KPI_CARD

    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = config.FILL_KPI_CARD
            left_s = Side(style='medium', color=config.COLOR_PRIMARY_NAVY) if c == c1 else Side(style='thin', color=config.COLOR_BORDER)
            right_s = Side(style='thin', color=config.COLOR_BORDER)
            top_s = Side(style='thin', color=config.COLOR_BORDER) if r == r1 else None
            bot_s = Side(style='thin', color=config.COLOR_BORDER) if r == r2 else None
            cell.border = Border(left=left_s, right=right_s, top=top_s, bottom=bot_s)

def draw_kpi_cards(ws, kpis, start_row=4):
    """
    Draws a horizontal strip of KPI cards (e.g., in executive dashboard tabs).
    Each card spans 2 columns with Row 4 = Label and Row 5 = Metric Value.
    kpis is a list of tuples: (label, val, note, is_alert)
    """
    for idx, (label, val, note, is_alert) in enumerate(kpis):
        col_start = 1 + (idx * 2)
        create_kpi_card(
            ws,
            col_start,
            start_row,
            label,
            val,
            is_alert=is_alert,
            label_on_top=True
        )
        ws.row_dimensions[start_row].height = 18
        ws.row_dimensions[start_row + 1].height = 28

def clean_ou_token(token):
    """Strips LDAP prefixes OU=, CN=, DC= from an Active Directory token."""
    if not token:
        return ""
    t = token.strip()
    for prefix in ["OU=", "CN=", "DC="]:
        if t.startswith(prefix):
            return t[len(prefix):].strip()
    return t
