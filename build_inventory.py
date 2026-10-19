"""
build_inventory.py
==================
All-in-one automation script to transform raw NinjaOne CSV exports into the
fully formatted, multi-tab executive 'Ninja Inventory' Excel workbook.

Requires only 2 CSV exports:
  1. Devices(1).csv  -> From NinjaOne: 'Export with all device details'
  2. Devices.csv     -> From NinjaOne: 'Export with custom fields' (for AD Install Date)
Note: Devices(2).csv ('Export visible columns') is 100% redundant and NOT needed!
"""

import os
import sys
import re
import csv
import glob
from datetime import datetime
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import config
from utils import parse_date as parse_ninja_date, safe_float, safe_int, normalize_hostname
from generate_dashboard_tabs import generate_dashboards
from build_ad_report import parse_ad_csv, build_ad_inventory_tab, build_ad_dashboard_tab
from build_huntress_report import parse_huntress_csv, build_huntress_inventory_tab, build_huntress_dashboard_tab

# Regex for volume parsing
re_cap = re.compile(r'\(([\d\.]+)\s*GiB\)')
re_bytes = re.compile(r'Capacity:\s*"(\d+)')
re_usage = re.compile(r'Usage %:\s*"(\d+)%"')

def build_inventory_sheet(csv_details_path, csv_custom_path, output_excel_path, csv_huntress_path=None, csv_ad_path=None):
    ad_dates = {}
    if csv_custom_path and os.path.exists(csv_custom_path):
        print(f"Reading custom fields from {csv_custom_path}...")
        with open(csv_custom_path, mode='r', encoding='utf-8', errors='ignore') as f:
            reader = csv.reader(f)
            header = [c.replace('\ufeff', '').strip() for c in next(reader)]
            dev_idx = header.index('Device') if 'Device' in header else -1
            if dev_idx == -1 and 'Display Name' in header:
                dev_idx = header.index('Display Name')
            if dev_idx == -1 and 'Name' in header:
                dev_idx = header.index('Name')
            ad_idx = header.index('AD Install Date') if 'AD Install Date' in header else -1
            if ad_idx == -1:
                for idx, col in enumerate(header):
                    if 'install date' in col.lower():
                        ad_idx = idx
                        break
            
            for row in reader:
                if row and dev_idx != -1 and ad_idx != -1 and len(row) > max(dev_idx, ad_idx):
                    dev_name = row[dev_idx].strip()
                    ad_val = row[ad_idx].strip()
                    if ad_val:
                        parsed = parse_ninja_date(ad_val)
                        if parsed:
                            ad_dates[dev_name.upper()] = parsed
                            ad_dates[normalize_hostname(dev_name)] = parsed

    print(f"Loaded {len(ad_dates)} AD Install Dates.")
    has_ad_data = len(ad_dates) > 0

    print(f"Reading device details from {csv_details_path}...")
    devices = []
    parsed_online_dates = []

    with open(csv_details_path, mode='r', encoding='utf-8', errors='ignore') as f:
        reader = csv.reader(f)
        header = [c.replace('\ufeff', '').strip() for c in next(reader)]
        # Map all headers dynamically
        col_map = {name: i for i, name in enumerate(header)}
        
        for row in reader:
            if not row or not any(row):
                continue
            def get_val(col_name):
                idx = col_map.get(col_name)
                return row[idx].strip() if idx is not None and idx < len(row) else ""

            dev_name = get_val('Display Name') or get_val('Device') or get_val('Device Name') or get_val('Name')
            if not dev_name:
                continue

            last_online_str = get_val('Last Online_formatted')
            if last_online_str:
                p_date = parse_ninja_date(last_online_str)
                if p_date:
                    parsed_online_dates.append(p_date)

            ram_val = safe_float(get_val('Memory Capacity GiB'))
            clock_raw = get_val('Processors Max Clock Speed')
            clock_val = None
            if clock_raw:
                c_num = safe_float(clock_raw)
                if c_num is not None:
                    clock_val = round(c_num / 1e9, 3) if c_num > 100000 else round(c_num, 3)

            devices.append({
                'Device': dev_name,
                'Location': get_val('Location'),
                'Device Role': get_val('Device Role'),
                'Device Model': get_val('Device Model'),
                'Serial Number': get_val('Serial Number'),
                'Last Login': get_val('Last Login'),
                'Last Uptime_formatted': last_online_str,
                'Last Boot Time_formatted': get_val('Last Boot Time_formatted'),
                'Policy': get_val('Policy'),
                'OS': get_val('OS Release Id'),
                'Last Update_formatted': get_val('Last Update_formatted'),
                'RAM': ram_val,
                'Volumes': get_val('Volumes'),
                'Processors Name': get_val('Processors Name'),
                'Processors Max Clock Speed': clock_val,
                'Warranty Start Date_formatted': get_val('Warranty Start Date_formatted'),
                'Warranty End Date_formatted': get_val('Warranty End Date_formatted')
            })

    # Robust reference date: max of today and any online timestamp in the file
    today_date = datetime.now().date()
    max_device_date = max(parsed_online_dates) if parsed_online_dates else today_date
    ref_date = max(today_date, max_device_date)
    print(f"Parsed {len(devices)} devices. Reference Audit Date: {ref_date}")

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ninja Inventory"

    headers = config.INVENTORY_HEADERS
    ws.append(headers)

    for d in devices:
        dev_name = d['Device']
        # Uptime days - clamped to >= 0
        up_date = parse_ninja_date(d['Last Uptime_formatted'])
        up_days = max(0, (ref_date - up_date).days) if up_date else None

        # Boot days - clamped to >= 0
        boot_date = parse_ninja_date(d['Last Boot Time_formatted'])
        boot_days = max(0, (ref_date - boot_date).days) if boot_date else None

        # Update days - clamped to >= 0
        upd_date = parse_ninja_date(d['Last Update_formatted'])
        upd_days = max(0, (ref_date - upd_date).days) if upd_date else None

        # Volume parsing
        vol_str = d['Volumes']
        cap_val = None
        free_val = None
        use_val = None
        if vol_str:
            cap_m = re_cap.search(vol_str)
            if cap_m:
                cap_val = float(cap_m.group(1))
            else:
                byte_m = re_bytes.search(vol_str)
                if byte_m:
                    cap_val = round(int(byte_m.group(1)) / (1024**3), 1)

            use_m = re_usage.search(vol_str)
            if use_m:
                use_val = float(use_m.group(1)) / 100.0

            if cap_val is not None and use_val is not None:
                free_val = round(cap_val * (1.0 - use_val), 1)

        # AD Install & Refresh Year
        ad_date = ad_dates.get(dev_name.upper()) or ad_dates.get(normalize_hostname(dev_name)) or ad_dates.get(dev_name)
        if ad_date:
            rf_year = ad_date.year + 3
        else:
            w_end_p = parse_ninja_date(d['Warranty End Date_formatted'])
            rf_year = (w_end_p.year + 3) if w_end_p else None

        # Warranty status & days
        w_end_date = parse_ninja_date(d['Warranty End Date_formatted'])
        w_days = (w_end_date - ref_date).days if w_end_date else None
        if w_days is not None:
            if w_days < 0:
                w_status = 'Expired'
            elif w_days <= 90:
                w_status = 'Expiring Soon'
            else:
                w_status = 'Active'
        else:
            w_status = 'Unknown'

        ws.append([
            dev_name,
            d['Location'],
            d['Device Role'],
            d['Device Model'],
            d['Serial Number'],
            d['Last Login'],
            d['Last Uptime_formatted'],
            up_days,
            d['Last Boot Time_formatted'],
            boot_days,
            d['Policy'],
            d['OS'],
            d['Last Update_formatted'],
            upd_days,
            d['RAM'],
            cap_val,
            free_val,
            use_val,
            d['Processors Name'],
            d['Processors Max Clock Speed'],
            ad_date,
            rf_year,
            d['Warranty Start Date_formatted'],
            d['Warranty End Date_formatted'],
            w_status,
            w_days
        ])

    print("Styling Inventory tab with full color tiers...")
    # Centralized Styling Palette from config
    font_header   = config.FONT_TBL_HEADER
    fill_header   = config.FILL_NAVY_HEADER
    font_data     = config.FONT_TBL_DATA
    fill_zebra    = config.FILL_ZEBRA
    fill_white    = config.FILL_WHITE
    thin_border   = config.BORDER_THIN

    fill_week     = config.FILL_WEEK
    font_week     = config.FONT_WEEK
    fill_month    = config.FILL_MONTH
    font_month    = config.FONT_MONTH
    fill_3months  = config.FILL_3MONTHS
    font_3months  = config.FONT_3MONTHS
    fill_6months  = config.FILL_6MONTHS
    font_6months  = config.FONT_6MONTHS
    fill_lavender = config.FILL_LAVENDER
    font_lavender = config.FONT_LAVENDER
    fill_active   = config.FILL_ACTIVE
    font_active   = config.FONT_ACTIVE
    fill_staging  = config.FILL_STAGING
    font_staging  = config.FONT_STAGING

    def color_maintenance_days(cell, days):
        if days is not None:
            try:
                d = float(days)
                if d <= 7:
                    pass  # Normal zebra
                elif d <= 30:
                    cell.fill = fill_week
                    cell.font = font_week
                elif d <= 90:
                    cell.fill = fill_month
                    cell.font = font_month
                elif d <= 180:
                    cell.fill = fill_3months
                    cell.font = font_3months
                elif d <= 365:
                    cell.fill = fill_6months
                    cell.font = font_6months
                else:
                    # > 365 days -> Critical Lavender
                    cell.fill = fill_lavender
                    cell.font = font_lavender
            except (ValueError, TypeError):
                pass

    def color_disk_usage(cell, val):
        if val is not None:
            try:
                pct = float(val)
                if pct < 0.60:
                    pass  # Healthy zebra
                elif pct < 0.70:
                    cell.fill = fill_week
                    cell.font = font_week
                elif pct < 0.75:
                    cell.fill = fill_month
                    cell.font = font_month
                elif pct < 0.85:
                    cell.fill = fill_3months
                    cell.font = font_3months
                elif pct < 0.90:
                    cell.fill = fill_6months
                    cell.font = font_6months
                else:
                    # >= 90% Emergency -> Critical Lavender
                    cell.fill = fill_lavender
                    cell.font = font_lavender
            except (ValueError, TypeError):
                pass

    # Header Row Formatting
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(1, col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = config.ALIGN_HEADER

    ws.row_dimensions[1].height = 28

    curr_year = ref_date.year

    for r in range(2, ws.max_row + 1):
        bg = fill_zebra if r % 2 == 0 else fill_white
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = font_data
            cell.fill = bg
            cell.border = thin_border
            cell.alignment = config.ALIGN_LEFT

        # Number formatting
        ws.cell(r, 8).number_format = '#,##0'    # Uptime Days
        ws.cell(r, 10).number_format = '#,##0'   # Boot Days
        ws.cell(r, 14).number_format = '#,##0'   # Update Days
        ws.cell(r, 15).number_format = '#,##0.0' # RAM
        ws.cell(r, 16).number_format = '#,##0.0' # Capacity
        ws.cell(r, 17).number_format = '#,##0.0' # Free
        ws.cell(r, 18).number_format = '0%'      # Usage %
        ws.cell(r, 20).number_format = '#,##0.000' # Clock Speed
        if ws.cell(r, 21).value:
            ws.cell(r, 21).number_format = 'YYYY-MM-DD'
        ws.cell(r, 22).number_format = '0'       # Refresh Year
        ws.cell(r, 26).number_format = '#,##0'   # Warranty Days

        # Align center
        for col_c in [8, 10, 12, 14, 21, 22, 25, 26]:
            ws.cell(r, col_c).alignment = config.ALIGN_CENTER

        # 1. Device Hostname highlight (Staging / Default names)
        dev_val = str(ws.cell(r, 1).value or "")
        if dev_val.startswith("WINDOWS-") or "TEMP" in dev_val:
            ws.cell(r, 1).fill = fill_staging
            ws.cell(r, 1).font = font_staging

        # 2. Last Uptime - Days (Col 8)
        color_maintenance_days(ws.cell(r, 8), ws.cell(r, 8).value)

        # 3. Last Boot Time - Days (Col 10)
        color_maintenance_days(ws.cell(r, 10), ws.cell(r, 10).value)

        # 4. Last Update - Days (Col 14)
        color_maintenance_days(ws.cell(r, 14), ws.cell(r, 14).value)

        # 5. Primary Drive Usage % (Col 18)
        color_disk_usage(ws.cell(r, 18), ws.cell(r, 18).value)

        # 6. Refresh Year (Col 22) - Distinguish computers meant to be replaced this year / overdue
        rf_val = ws.cell(r, 22).value
        if rf_val is not None:
            try:
                rf_int = int(rf_val)
                if rf_int < curr_year:
                    # Overdue from past years -> Critical Lavender
                    ws.cell(r, 22).fill = fill_lavender
                    ws.cell(r, 22).font = font_lavender
                elif rf_int == curr_year:
                    # Due this year -> Peach / Alert
                    ws.cell(r, 22).fill = fill_month
                    ws.cell(r, 22).font = font_3months
            except (ValueError, TypeError):
                pass

        # 7. Warranty Status (Col 25)
        st_val = str(ws.cell(r, 25).value or "")
        if st_val == "Active":
            ws.cell(r, 25).fill = fill_active
            ws.cell(r, 25).font = font_active
        elif st_val == "Expiring Soon":
            ws.cell(r, 25).fill = fill_week
            ws.cell(r, 25).font = font_week
        elif st_val == "Expired":
            ws.cell(r, 25).fill = fill_month
            ws.cell(r, 25).font = font_3months

        # 8. Warranty Days Remaining (Col 26)
        wd_val = ws.cell(r, 26).value
        if wd_val is not None:
            try:
                wd_num = float(wd_val)
                if wd_num > 90:
                    pass  # Normal active zebra
                elif wd_num >= 0:
                    # Expiring Soon (0 - 90 days) -> Soft Yellow
                    ws.cell(r, 26).fill = fill_week
                    ws.cell(r, 26).font = font_week
                elif wd_num >= -365:
                    # Expired within a year -> Soft Peach
                    ws.cell(r, 26).fill = fill_month
                    ws.cell(r, 26).font = font_3months
                else:
                    # Expired > 1 year ago -> Critical Lavender
                    ws.cell(r, 26).fill = fill_lavender
                    ws.cell(r, 26).font = font_lavender
            except (ValueError, TypeError):
                pass

    # Column Widths
    for c_idx, width in config.INVENTORY_COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    # Convert Master Inventory into an official Excel Table with auto-filters and sorting (if data rows exist)
    if ws.max_row > 1:
        tbl_inv = Table(displayName="Table_MasterInventory", ref=f"A1:Z{ws.max_row}")
        tbl_inv.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight8",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        ws.add_table(tbl_inv)

    # Freeze row 1 + Columns A & B (Device & Location pinned)
    ws.freeze_panes = "C2"

    # -------------------------------------------------------------------------
    # Optional: Build 'AD Inventory' and 'AD Dashboard' tabs
    # -------------------------------------------------------------------------
    ad_records = None
    if csv_ad_path and os.path.exists(csv_ad_path):
        print(f"Reading Active Directory raw export from {csv_ad_path}...")
        ad_records = parse_ad_csv(csv_ad_path)
        build_ad_inventory_tab(wb, ad_records, title="AD Inventory")
        build_ad_dashboard_tab(wb, ad_records, title="AD Dashboard")

    # -------------------------------------------------------------------------
    # Optional: Build 'Huntress Inventory' and 'Huntress Dashboard' tabs
    # -------------------------------------------------------------------------
    has_huntress = False
    if csv_huntress_path and os.path.exists(csv_huntress_path):
        print(f"Reading Huntress agents export from {csv_huntress_path}...")
        hunt_agents = parse_huntress_csv(csv_huntress_path)
        build_huntress_inventory_tab(wb, hunt_agents, title="Huntress Inventory")
        build_huntress_dashboard_tab(wb, hunt_agents, title="Huntress Dashboard")
        has_huntress = True

    # Generate executive dashboard tabs in-process (direct function call)
    print("Generating Executive Dashboard tabs...")
    generate_dashboards(
        wb,
        huntress_csv=csv_huntress_path if has_huntress else None,
        skip_lifecycle=not has_ad_data
    )

    # Enforce exact desired tab order
    desired_order = [
        "Ninja Inventory",
        "AD Inventory",
        "AD Computers",
        "Huntress Inventory",
        "Huntress Agents",
        "Patch & Reboot Health",
        "Lifecycle",
        "Performance & Capacity",
        "Site Allocation",
        "Stale Assets",
        "AD Dashboard",
        "AD Cross ref",
        "Huntress Dashboard",
        "Huntress Cross ref",
        "Inventory",
        "Dashboard",
        "Active Directory",
        "Huntress Endpoint Protection",
        "Executive Summary"
    ]
    ordered_sheets = [wb[s] for s in desired_order if s in wb.sheetnames]
    for s in wb.sheetnames:
        if wb[s] not in ordered_sheets:
            ordered_sheets.append(wb[s])
    wb._sheets = ordered_sheets

    if "Ninja Inventory" in wb.sheetnames:
        wb.active = wb["Ninja Inventory"]
    elif "Inventory" in wb.sheetnames:
        wb.active = wb["Inventory"]

    out_dir = os.path.dirname(os.path.abspath(output_excel_path))
    os.makedirs(out_dir, exist_ok=True)

    try:
        wb.save(output_excel_path)
        print(f"Workbook successfully saved to: {output_excel_path}")
    except PermissionError:
        raise PermissionError(
            f"Cannot save '{os.path.basename(output_excel_path)}' because the file is currently open in Microsoft Excel or another program.\n"
            "Please close the workbook in Excel and re-run the report."
        )

    # -------------------------------------------------------------------------
    # Optional: Automatically Export Actionable Flagged Devices to Google Tasks CSV
    # -------------------------------------------------------------------------
    if csv_huntress_path and csv_ad_path:
        try:
            from task_exporter import export_flagged_devices_tasks
            # Route tasks to reports/Tasks/
            out_dir = os.path.dirname(os.path.abspath(output_excel_path))
            if os.path.basename(out_dir).lower() == "extensive reports":
                tasks_dir = os.path.join(os.path.dirname(out_dir), "Tasks")
            else:
                tasks_dir = os.path.join("reports", "Tasks")
            tasks_dir = os.path.normpath(tasks_dir)

            print(f"Generating actionable Google Tasks export in '{tasks_dir}'...")
            t_summary = export_flagged_devices_tasks(wb, output_dir=tasks_dir, ref_date=ref_date)
            print(f"SUCCESS: Exported {t_summary['total_flagged']} flagged devices across {t_summary['total_clinics']} clinics.")
            print(f"Master Tasks File: {t_summary['master_file']}")
        except Exception as e:
            print(f"Notice: Could not generate tasks export: {e}")

if __name__ == "__main__":
    from detect_files import find_best_ninja_csvs

    today_str = datetime.now().strftime("%m%d%Y")
    
    # Smart Auto-Detection
    detected = find_best_ninja_csvs()
    
    # Parse positional or named arguments
    # Usage: python build_inventory.py [details.csv] [custom.csv] [output.xlsx] [huntress.csv] [ad.csv]
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    huntress_cli = None
    ad_cli = None
    for i, a in enumerate(sys.argv):
        if a == "--huntress" and i + 1 < len(sys.argv):
            huntress_cli = sys.argv[i + 1]
        elif a == "--ad" and i + 1 < len(sys.argv):
            ad_cli = sys.argv[i + 1]

    d1 = args[0] if len(args) > 0 else detected.get("details_file")
    d0 = args[1] if len(args) > 1 else detected.get("custom_file")
    out_file = args[2] if len(args) > 2 else f"Ninja Inventory {today_str}.xlsx"
    if "--no-huntress" in sys.argv:
        dh = None
    else:
        dh = huntress_cli or (args[3] if len(args) > 3 else detected.get("huntress_file"))

    dad = ad_cli or (args[4] if len(args) > 4 else None)

    print(f"--- Starting Automated Ninja Inventory Pipeline ---")
    print(f"Details CSV:       {d1} (auto-detected: {bool(not args or len(args) < 1)})")
    print(f"Custom Fields CSV: {d0} (auto-detected: {bool(not args or len(args) < 2)})")
    print(f"Huntress CSV:      {dh or 'None (Optional)'} (auto-detected: {bool(not huntress_cli and (len(args) < 4))})")
    print(f"AD Raw CSV:        {dad or 'None (Optional)'}")
    print(f"Output Workbook:   {out_file}")
    if detected.get("redundant_found"):
        print(f"Ignored Files:     {[os.path.basename(f) for f in detected['redundant_found']]} (redundant)")
    print(f"---------------------------------------------------")
    
    if not d1 or not os.path.exists(d1):
        print(f"ERROR: Could not locate a valid Device Details CSV export.")
        sys.exit(1)

    build_inventory_sheet(d1, d0, out_file, csv_huntress_path=dh, csv_ad_path=dad)
    print(f"DONE! Workbook ready at: {out_file}")
